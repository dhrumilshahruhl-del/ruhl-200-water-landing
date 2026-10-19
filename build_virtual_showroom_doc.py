from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "200-Water-Virtual-Showroom-Technical-Overview.docx"
ROOF_IMG = ROOT / "scripts" / "roof-detail-check.png"

BLUE = "00375A"
BLACK = "000000"
WHITE = "FFFFFF"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def run(paragraph, text, *, size=10, bold=False, color=BLACK, font="Arial"):
    r = paragraph.add_run(text)
    r.font.name = font
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    r.font.size = Pt(max(size, 10))
    r.font.color.rgb = rgb(color)
    r.bold = bold
    return r


def set_cell_border(cell, color="000000", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def shade_white(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), WHITE)


def set_table_margins(table, top=80, bottom=80, start=100, end=100):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=10640):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")


def setup_doc(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = rgb(BLACK)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.08


def paragraph(doc, text="", *, after=4, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align:
        p.alignment = align
    if text:
        run(p, text)
    return p


def heading(doc, text, *, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run(p, text, size=13 if level == 1 else 11, bold=True, color=BLUE)
    return p


def bullet(doc, text, *, after=2):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.05
    run(p, "- ", bold=True)
    run(p, text)
    return p


def table_text(cell, label, text):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    shade_white(cell)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, label, bold=True)
    if text:
        run(p, "\n" + text)


def table_body(cell, text):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    shade_white(cell)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run(p, text)


def code_line(doc, label, code):
    p = paragraph(doc, after=4, before=2)
    run(p, label + ": ", bold=True)
    run(p, code, font="Consolas", size=10)


def code_box(doc, label, code):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    set_table_width(table)
    set_table_margins(table, top=90, bottom=90, start=110, end=110)
    cell = table.cell(0, 0)
    shade_white(cell)
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run(p, label + "\n", bold=True)
    run(p, code, font="Consolas", size=10)


def title(doc):
    p = paragraph(doc, after=2)
    run(p, "200 Water Rooftop Virtual Showroom", size=18, bold=True, color=BLUE)
    p = paragraph(
        doc,
        "Technical overview of the interactive 3D rooftop experience, based strictly on project files.",
        after=7,
    )


def overview(doc):
    heading(doc, "Purpose & Experience")
    bullet(doc, "Creates a scroll-driven 3D building tour that transitions into a roof-only inspection mode for rooftop systems.")
    bullet(doc, "Users open the detail view, rotate the roof, and inspect callouts, leader lines, legend, and flow overlays.")

    heading(doc, "Technology Stack")
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    set_table_width(table)
    set_table_margins(table)
    rows = [
        ("Runtime", "Static HTML/CSS/JavaScript; Tailwind CDN; Lucide; Roboto; no build step documented."),
        ("3D + motion", "Three.js 0.170, GLTFLoader, MeshoptDecoder, WebGLRenderer, PerspectiveCamera, GSAP ScrollTrigger."),
        ("Models", "Full building: logos/200_water.opt.glb. Roof detail: logos/200 water Roof Only.glb."),
    ]
    for i, (label, text) in enumerate(rows):
        table_text(table.cell(i, 0), label, "")
        table_body(table.cell(i, 1), text)


def engineering(doc):
    heading(doc, "Engineering Workflow")
    bullet(doc, "model-preload.js centralizes GLB loading: building first, roof deferred and prefetched later.")
    bullet(doc, "The building GLB is cloned, material-prepared, centered, uniformly scaled, rotated, and framed with a custom camera/light rig.")
    bullet(doc, "Rooftop framing samples upper-roof vertices and adjusts projection until the inspection view fits cleanly.")
    bullet(doc, "Roof-only alignment replays the building transform, then screen-aligns the separate roof model.")
    code_box(
        doc,
        "Code evidence: deferred model loading",
        "export const buildingGltfPromise = preloadGltf(paths.building);\n"
        "if (!roofGltfPromise) roofGltfPromise = preloadGltf(paths.roof);",
    )

    heading(doc, "Interaction Engineering")
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    set_table_width(table)
    set_table_margins(table)
    rows = [
        ("Scroll choreography", "A smoothed playhead synchronizes camera, FOV, lights, CTA state, render timing, and roof prefetch."),
        ("DOM + WebGL sync", "rooftop-labels.js projects 3D anchors so DOM labels, leader lines, legend, and SVG flow paths stay aligned with the model."),
    ]
    for i, (label, text) in enumerate(rows):
        table_text(table.cell(i, 0), label, "")
        table_body(table.cell(i, 1), text)

    heading(doc, "Performance Engineering")
    bullet(doc, "Primary GLB is preloaded and cached; roof loading is deferred until needed.")
    bullet(doc, "MeshoptDecoder handles compressed GLBs; optimized metadata includes EXT_meshopt_compression/KHR_mesh_quantization.")
    bullet(doc, "Renderer DPR is capped, mobile antialiasing is disabled, and rendering pauses outside the active 3D section.")

    heading(doc, "Architecture")
    p = paragraph(doc, after=2)
    run(
        p,
        "User scroll / CTA -> GSAP ScrollTrigger + smoothing -> Three.js scene/camera/lights -> full-building GLB and roof-only alignment -> DOM/SVG labels and flow overlays",
        bold=True,
    )


def page_two(doc):
    doc.add_page_break()
    heading(doc, "Rooftop Detail View")
    if ROOF_IMG.exists():
        visual = doc.add_table(rows=1, cols=2)
        visual.autofit = False
        set_table_width(visual)
        set_table_margins(visual, top=80, bottom=80, start=90, end=90)
        left = visual.cell(0, 0)
        right = visual.cell(0, 1)
        for cell in (left, right):
            shade_white(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        left.paragraphs[0].add_run().add_picture(str(ROOF_IMG), width=Inches(4.15))
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(2)
        run(rp, "Rooftop interaction shown", bold=True)
        for item in [
            "Roof-only inspection state loaded from the project rooftop GLB.",
            "Equipment labels, leader lines, legend, and flow overlays presented in one coordinated view.",
            "Screen-space DOM/SVG annotations aligned to the underlying WebGL model.",
        ]:
            p = right.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.first_line_indent = Inches(-0.15)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.05
            run(p, "- ")
            run(p, item)
        p = paragraph(
            doc,
            "Figure: active rooftop detail view captured from the project implementation.",
            after=5,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    heading(doc, "Most Complex Implementation Areas")
    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    set_table_width(table)
    set_table_margins(table)
    rows = [
        ("Cross-model alignment", "Transform replay plus projected-bounds matching bridges full-building and roof-only GLBs with different coordinate assumptions."),
        ("Rooftop camera math", "Roof vertex sampling, projection checks, and iterative distance fitting create the inspection composition."),
        ("Overlay precision", "DOM labels and SVG flows remain attached to live 3D anchors through resize, drag, and model swapping."),
    ]
    for i, (label, text) in enumerate(rows):
        table_text(table.cell(i, 0), label, "")
        table_body(table.cell(i, 1), text)

    code_box(
        doc,
        "Code evidence: scroll smoothing",
        "const delta = targetProgress - playhead.progress;\n"
        "playhead.progress += delta * smoothing;\n"
        "timeline.progress(playhead.progress);",
    )

    heading(doc, "Evidence & Limits")
    paragraph(
        doc,
        "Evidence reviewed: runtime HTML/CSS/JS modules, README, building-scroll-recipe.md, wrangler.jsonc, audit scripts/screenshots, and GLB metadata from logos/*.glb.",
        after=2,
    )
    bullet(doc, "Blender is verified only as source/export evidence; active optimized GLBs report glTF-Transform v4.4.2.")
    bullet(doc, "Not verified: source .blend file, backend code, package/build pipeline, or live import/use of rooftop-retrofit.js.")


def build():
    doc = Document()
    setup_doc(doc)
    title(doc)
    overview(doc)
    engineering(doc)
    page_two(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
